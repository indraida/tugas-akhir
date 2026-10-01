from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from modules.tk_report import filter_tk_report_source, generate_tk_excel, generate_tk_filename, generate_tk_pdf, get_report_period_months, prepare_tk_report_data


def sample_data():
    return pd.DataFrame([
        {"NIP": "198001010000000001", "Nama Pegawai": "Pegawai A", "Unit Kerja": "BKAD", "Tahun": 2026, "Bulan": "Januari", "TK": 2},
        {"NIP": "198001010000000001", "Nama Pegawai": "Pegawai A", "Unit Kerja": "BKAD", "Tahun": 2026, "Bulan": "Februari", "TK": 1},
        {"NIP": "198002020000000002", "Nama Pegawai": "Pegawai B", "Unit Kerja": "BKAD", "Tahun": 2026, "Bulan": "Februari", "TK": 2},
        {"NIP": "198003030000000003", "Nama Pegawai": "Pegawai C", "Unit Kerja": "BKD", "Tahun": 2026, "Bulan": "Januari", "TK": 3},
        {"NIP": "198004040000000004", "Nama Pegawai": "Pegawai D", "Unit Kerja": "Dinkes", "Tahun": 2026, "Bulan": "Januari", "TK": 0},
    ])


def test_report_period_mapping():
    assert get_report_period_months("TRIWULAN_I") == [1, 2, 3]
    assert get_report_period_months("TRIWULAN_II") == [4, 5, 6]
    assert get_report_period_months("TRIWULAN_III") == [7, 8, 9]
    assert get_report_period_months("TRIWULAN_IV") == [10, 11, 12]
    assert get_report_period_months("TAHUNAN") == list(range(1, 13))


def test_shared_scope_produces_identical_report_for_analysis_and_report_page():
    source = sample_data().assign(**{"Jenis Pegawai": ["PNS", "PNS", "PPPK", "PPPK", "PNS"]})
    analysis_source = filter_tk_report_source(source, 2026, "Semua OPD", "Semua Jenis Pegawai")
    page_source = filter_tk_report_source(source, 2026, "Semua OPD", "Semua Jenis Pegawai")
    analysis_report = prepare_tk_report_data(analysis_source, "Semua Jenis Pegawai", "TRIWULAN_I")
    page_report = prepare_tk_report_data(page_source, "Semua Jenis Pegawai", "TRIWULAN_I")
    keys = ["total_opd", "total_employees", "grand_total_tk", "opds"]
    assert {key: analysis_report[key] for key in keys} == {key: page_report[key] for key in keys}


@pytest.mark.parametrize("period", ["TRIWULAN_I", "TRIWULAN_II", "TRIWULAN_III", "TRIWULAN_IV", "TAHUNAN"])
@pytest.mark.parametrize("employee_type", ["Semua Jenis Pegawai", "PNS", "PPPK"])
def test_report_contract_accepts_all_periods_opd_and_employee_types(period, employee_type):
    source = sample_data().assign(**{"Jenis Pegawai": ["PNS", "PNS", "PPPK", "PPPK", "PNS"]})
    scoped = source[source["Unit Kerja"].eq("BKAD")]
    report = prepare_tk_report_data(scoped, employee_type, period)
    assert report["period_type"] == period
    assert report["months"] == list(range(1, 13))


def test_report_model_keeps_zero_tk_opd_and_uses_dynamic_title_period():
    report = prepare_tk_report_data(sample_data(), "PNS")
    assert report["employee_type"] == "PNS"
    assert "DAFTAR NAMA PNS" in report["title"]
    assert report["period_text"] == "PERIODE JANUARI 2026 s.d. DESEMBER 2026"
    assert [opd["opd_name"] for opd in report["opds"]] == ["BKAD", "BKD", "Dinkes"]
    dinkes = next(opd for opd in report["opds"] if opd["opd_name"] == "Dinkes")
    assert dinkes["total_tk"] == 0
    assert dinkes["employees"] == []
    assert report["grand_total"] == 8
    assert all(employee["total_tk"] > 0 for opd in report["opds"] for employee in opd["employees"])


def test_all_employee_types_use_asn_title_and_quarter_period():
    january = sample_data().query("Bulan == 'Januari'")
    report = prepare_tk_report_data(january, "Semua Jenis Pegawai", "TRIWULAN_I")
    assert "DAFTAR NAMA ASN" in report["title"]
    assert report["period_text"] == "PERIODE JANUARI 2026 s.d. MARET 2026"
    assert report["active_months"] == [1, 2, 3]
    assert report["months"] == list(range(1, 13))


def test_cross_period_totals_and_month_columns_are_isolated():
    source = pd.concat([sample_data(), pd.DataFrame([{"NIP": "198001010000000001", "Nama Pegawai": "Pegawai A", "Unit Kerja": "BKAD", "Tahun": 2026, "Bulan": "April", "TK": 5}])], ignore_index=True)
    first = prepare_tk_report_data(source, period_type="TRIWULAN_I")
    second = prepare_tk_report_data(source, period_type="TRIWULAN_II")
    annual = prepare_tk_report_data(source, period_type="TAHUNAN")
    assert first["active_months"] == [1, 2, 3] and first["grand_total"] == 8
    assert second["active_months"] == [4, 5, 6] and second["grand_total"] == 5
    assert annual["active_months"] == list(range(1, 13)) and annual["grand_total"] == 13
    assert first["months"] == second["months"] == annual["months"] == list(range(1, 13))
    quarter_book = load_workbook(BytesIO(generate_tk_excel(first)), data_only=True)["REKAP TK"]
    annual_book = load_workbook(BytesIO(generate_tk_excel(annual)), data_only=True)["REKAP TK"]
    assert quarter_book.max_column == 20
    assert annual_book.max_column == 20
    assert [quarter_book.cell(12, column).value for column in range(7, 19)] == ["JAN", "FEB", "MAR", "APR", "MEI", "JUN", "JUL", "AGT", "SEP", "OKT", "NOV", "DES"]
    employee_row = next(cell.row for cell in quarter_book["F"] if cell.value == "198001010000000001")
    assert [quarter_book.cell(employee_row, column).value for column in range(7, 19)] == [2, 1, "-", "-", "-", "-", "-", "-", "-", "-", "-", "-"]
    assert quarter_book.cell(employee_row, 19).value == 3
    assert generate_tk_filename(first, "xlsx") == "Laporan_TK_Triwulan_I_2026.xlsx"
    assert generate_tk_filename(first, "pdf", "BKD") == "Laporan_TK_BKD_Triwulan_I_2026.pdf"


def test_excel_and_pdf_share_total_and_excel_keeps_nip_as_text():
    report = prepare_tk_report_data(sample_data(), "PPPK")
    excel = generate_tk_excel(report)
    pdf = generate_tk_pdf(report)
    workbook = load_workbook(BytesIO(excel), data_only=True)
    assert workbook.sheetnames == ["REKAP TK"]
    assert workbook["REKAP TK"].page_setup.orientation == "landscape"
    nip_cells = [cell for row in workbook["REKAP TK"].iter_rows() for cell in row if cell.value == "198001010000000001"]
    assert nip_cells and nip_cells[0].number_format == "@"
    employee_totals = [employee["total_tk"] for opd in report["opds"] for employee in opd["employees"]]
    assert sum(employee_totals) == report["grand_total"] == 8
    assert excel.startswith(b"PK")
    assert pdf.startswith(b"%PDF")
