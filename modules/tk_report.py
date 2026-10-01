"""Generator laporan TK in-memory tanpa file template eksternal."""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


MONTHS = [
    "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember",
]
MONTH_SHORT = ["JAN", "FEB", "MAR", "APR", "MEI", "JUN", "JUL", "AGT", "SEP", "OKT", "NOV", "DES"]
REPORT_PERIODS = {
    "TRIWULAN_I": [1, 2, 3], "TRIWULAN_II": [4, 5, 6],
    "TRIWULAN_III": [7, 8, 9], "TRIWULAN_IV": [10, 11, 12],
    "TAHUNAN": list(range(1, 13)),
}
PERIOD_NAMES = {
    "TRIWULAN_I": "Triwulan I", "TRIWULAN_II": "Triwulan II",
    "TRIWULAN_III": "Triwulan III", "TRIWULAN_IV": "Triwulan IV",
    "TAHUNAN": "Tahunan",
}


def get_report_period_months(period_type: str) -> list[int]:
    if period_type not in REPORT_PERIODS:
        raise ValueError(f"Periode laporan tidak dikenal: {period_type}")
    return REPORT_PERIODS[period_type].copy()


def get_report_period_label(period_type: str, year_text: str) -> tuple[str, str]:
    months = get_report_period_months(period_type)
    period_name = "TAHUN" if period_type == "TAHUNAN" else PERIOD_NAMES[period_type].upper()
    heading = f"{period_name} {year_text}" if period_type == "TAHUNAN" else f"{period_name} TAHUN {year_text}"
    return heading, f"PERIODE {MONTHS[months[0] - 1].upper()} {year_text} s.d. {MONTHS[months[-1] - 1].upper()} {year_text}"


def _employee_label(employee_type: str) -> str:
    return employee_type if employee_type in {"PNS", "PPPK"} else "ASN"


def _period_text(years: list[int], visible_months: list[str]) -> str:
    if not years or not visible_months:
        return "PERIODE AKTIF"
    first_month, last_month = visible_months[0].upper(), visible_months[-1].upper()
    if len(years) == 1:
        suffix = str(years[0])
        return f"PERIODE {first_month} {suffix}" if len(visible_months) == 1 else f"PERIODE {first_month} {suffix} s.d. {last_month} {suffix}"
    return f"PERIODE {first_month} {years[0]} s.d. {last_month} {years[-1]}"


def filter_tk_report_source(
    source: pd.DataFrame,
    year: int | str | None = None,
    opd: str = "Semua OPD",
    employee_type: str = "Semua Jenis Pegawai",
) -> pd.DataFrame:
    """Scope bersama untuk semua pemakai laporan TK, tanpa agregasi baru."""
    scoped = source.copy()
    if year not in (None, "Semua Tahun"):
        scoped = scoped[pd.to_numeric(scoped["Tahun"], errors="coerce").eq(pd.to_numeric(year, errors="coerce"))]
    if opd != "Semua OPD":
        scoped = scoped[scoped["Unit Kerja"].astype(str).eq(str(opd))]
    if employee_type != "Semua Jenis Pegawai" and "Jenis Pegawai" in scoped.columns:
        scoped = scoped[scoped["Jenis Pegawai"].astype(str).eq(employee_type)]
    return scoped


def prepare_tk_report_data(
    source: pd.DataFrame,
    employee_type: str = "Semua Jenis Pegawai",
    period_type: str = "TAHUNAN",
    year: int | str | None = None,
    opd: str = "Semua OPD",
) -> dict[str, Any]:
    """Bentuk satu model laporan; setiap pegawai diidentifikasi dengan NIP."""
    required = {"NIP", "Nama Pegawai", "Unit Kerja", "Tahun", "Bulan", "TK"}
    missing = required.difference(source.columns)
    if missing:
        raise ValueError(f"Kolom laporan TK tidak tersedia: {', '.join(sorted(missing))}")

    scoped = filter_tk_report_source(source, year, opd, employee_type)
    scoped["NIP"] = scoped["NIP"].astype(str).str.strip()
    scoped["TK"] = pd.to_numeric(scoped["TK"], errors="coerce").fillna(0).astype(int)
    scoped["Tahun"] = pd.to_numeric(scoped["Tahun"], errors="coerce")
    years = sorted(scoped["Tahun"].dropna().astype(int).unique().tolist())
    report_opd_names = sorted(scoped["Unit Kerja"].dropna().astype(str).unique().tolist())
    selected_month_numbers = get_report_period_months(period_type)
    selected_month_names = [MONTHS[number - 1] for number in selected_month_numbers]
    scoped = scoped[scoped["Bulan"].astype(str).isin(selected_month_names)].copy()
    grouped = scoped.groupby(
        ["NIP", "Unit Kerja", "Tahun", "Bulan"], as_index=False, dropna=False
    ).agg(Nama=("Nama Pegawai", "first"), TK=("TK", "sum"))
    grouped = grouped[grouped["TK"].gt(0)].copy()
    year_text = str(years[0]) if len(years) == 1 else f"{years[0]}–{years[-1]}" if years else "AKTIF"
    period_label, period_text = get_report_period_label(period_type, year_text)
    label = _employee_label(employee_type)
    report: dict[str, Any] = {
        "title": f"DAFTAR NAMA {label} YANG TIDAK MASUK KERJA\nTANPA KETERANGAN MENURUT PERANGKAT DAERAH/UNIT KERJA",
        "year": years[0] if len(years) == 1 else year_text,
        "period_type": period_type,
        "period_name": PERIOD_NAMES[period_type],
        "period_label": period_label,
        "period_text": period_text,
        "active_months": selected_month_numbers,
        "months": list(range(1, 13)),
        "employee_type": label,
        "opds": [],
        "grand_total": 0,
    }
    # Seluruh OPD dalam scope laporan tetap ditampilkan, termasuk yang tidak
    # memiliki pegawai dengan TK pada periode aktif.
    for opd_name in report_opd_names:
        opd_rows = grouped[grouped["Unit Kerja"].astype(str) == opd_name]
        employees = []
        for nip, employee_rows in opd_rows.groupby("NIP", sort=True):
            month_values = employee_rows.groupby("Bulan")["TK"].sum()
            months = {index + 1: int(month_values.get(month, 0)) for index, month in enumerate(MONTHS)}
            total = sum(months.values())
            if total <= 0:
                continue
            employees.append({
                "nip": str(nip), "nama": str(employee_rows.iloc[0]["Nama"]),
                "months": months, "monthly_tk": months, "total_tk": total, "keterangan": "",
            })
        opd_total = sum(item["total_tk"] for item in employees)
        report["opds"].append({"opd": str(opd_name), "opd_name": str(opd_name), "employees": employees, "total_tk": opd_total})
        report["grand_total"] += opd_total
    report["total_opd"] = len(report["opds"])
    report["total_employees"] = len({employee["nip"] for opd in report["opds"] for employee in opd["employees"]})
    report["grand_total_tk"] = report["grand_total"]
    return report


def generate_tk_filename(report_data: dict[str, Any], extension: str, opd: str = "Semua OPD") -> str:
    scopes: list[str] = []
    if opd != "Semua OPD":
        scopes.append(re.sub(r"[^A-Za-z0-9]+", "_", opd).strip("_"))
    if report_data.get("employee_type") in {"PNS", "PPPK"}:
        scopes.append(str(report_data["employee_type"]))
    scope = "_".join(scopes) + ("_" if scopes else "")
    period = report_data["period_name"].replace(" ", "_")
    year = re.sub(r"[^0-9-]+", "", str(report_data["year"]))
    return f"Laporan_TK_{scope}{period}_{year}.{extension.lstrip('.')}"


def generate_tk_excel(report_data: dict[str, Any]) -> bytes:
    """Bangun workbook formal 12 bulan; periode hanya menentukan nilai aktif."""
    wb = Workbook()
    ws = wb.active
    ws.title = "REKAP TK"
    all_months = list(range(1, 13))
    total_column, notes_column, last_column = 19, 20, "T"
    ws.merge_cells("A1:T1"); ws["A1"] = "Lampiran Surat Sekretaris Daerah Provinsi Kalimantan Barat"
    ws.merge_cells("A2:T2"); ws["A2"] = "Nomor  : __________________"
    ws.merge_cells("A3:T3"); ws["A3"] = "Tanggal: __________________"
    ws.merge_cells("A5:T5"); ws["A5"] = report_data["title"].split("\n")[0]
    ws.merge_cells("A6:T6"); ws["A6"] = report_data["title"].split("\n")[1]
    ws.merge_cells("A7:T7"); ws["A7"] = report_data["period_label"]
    ws.merge_cells("A8:T8"); ws["A8"] = report_data["period_text"]
    for cell in (ws["A1"], ws["A2"], ws["A3"]):
        cell.font = Font(name="Arial", size=9)
        cell.alignment = Alignment(horizontal="left", vertical="center")
    for cell in (ws["A5"], ws["A6"], ws["A7"], ws["A8"]):
        cell.font = Font(name="Arial", size=12 if cell.row < 7 else 10, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for column, value in {"A": "NO", "B": "NAMA UNIT KERJA", "D": "NAMA", "F": "NIP", get_column_letter(total_column): "TOTAL", last_column: "KETERANGAN"}.items():
        ws[f"{column}10"] = value
    ws.merge_cells("B10:C12"); ws.merge_cells("D10:E12")
    for column in ["A", "F", "S", "T"]: ws.merge_cells(f"{column}10:{column}12")
    ws.merge_cells("G10:R10"); ws["G10"] = "JUMLAH KETIDAKHADIRAN MASUK KERJA TANPA KETERANGAN (TK)"
    for offset, quarter in [(0, "TRIWULAN I"), (3, "TRIWULAN II"), (6, "TRIWULAN III"), (9, "TRIWULAN IV")]:
        start, end = get_column_letter(7 + offset), get_column_letter(9 + offset)
        ws.merge_cells(f"{start}11:{end}11"); ws[f"{start}11"] = quarter
    for index, month_number in enumerate(all_months, start=7): ws.cell(12, index, MONTH_SHORT[month_number - 1])

    thin = Side(style="thin", color="808080")
    header_fill = PatternFill("solid", fgColor="EAF0F7")
    for row in ws.iter_rows(min_row=10, max_row=12, min_col=1, max_col=notes_column):
        for cell in row:
            cell.font = Font(name="Arial", size=9, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.fill = header_fill; cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    row_number = 13
    for opd_number, opd in enumerate(report_data["opds"], 1):
        ws.cell(row_number, 1, opd_number); ws.merge_cells(start_row=row_number, start_column=2, end_row=row_number, end_column=notes_column)
        ws.cell(row_number, 2, f"{opd['opd_name'].upper()} (TOTAL TK: {opd['total_tk']})")
        for cell in ws[row_number]:
            cell.font = Font(name="Arial", size=9, bold=True); cell.fill = PatternFill("solid", fgColor="F3F4F6")
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin); cell.alignment = Alignment(vertical="center")
        row_number += 1
        for employee_number, employee in enumerate(opd["employees"], 1):
            values = [employee_number, "", "", employee["nama"], "", employee["nip"]]
            values += [employee["monthly_tk"][month] or "-" for month in all_months]
            values += [employee["total_tk"], employee["keterangan"]]
            for column, value in enumerate(values, 1):
                cell = ws.cell(row_number, column, value)
                cell.font = Font(name="Arial", size=9); cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                cell.alignment = Alignment(horizontal="left" if column in {2, 3, 4, 5, 20} else "center", vertical="center", wrap_text=True)
            ws.cell(row_number, 6).number_format = "@"
            row_number += 1

    widths = {"A": 5, "B": 17, "C": 4, "D": 23, "E": 4, "F": 20, get_column_letter(total_column): 8, last_column: 18}
    for column in range(7, total_column): widths[get_column_letter(column)] = 5
    for column, width in widths.items(): ws.column_dimensions[column].width = width
    ws.row_dimensions[10].height = 34; ws.freeze_panes = "G13"; ws.print_title_rows = "10:12"
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = .25; ws.page_margins.right = .25; ws.page_margins.top = .4; ws.page_margins.bottom = .4
    ws.print_options.horizontalCentered = True; ws.print_area = f"A1:{last_column}{max(row_number - 1, 12)}"
    output = BytesIO(); wb.save(output)
    return output.getvalue()


def generate_tk_pdf(report_data: dict[str, Any]) -> bytes:
    """Bangun PDF A4 landscape dengan header tabel berulang."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
    except ImportError as exc:  # pragma: no cover - bergantung instalasi deployment
        raise RuntimeError("Dependency reportlab belum terpasang") from exc

    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=7 * mm, rightMargin=7 * mm, topMargin=8 * mm, bottomMargin=8 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TKTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=10, leading=13, alignment=TA_CENTER, spaceAfter=3 * mm)
    attachment_style = ParagraphStyle("TKAttachment", parent=styles["Normal"], fontName="Helvetica", fontSize=7.5, leading=10)
    story = [Paragraph("Lampiran Surat Sekretaris Daerah Provinsi Kalimantan Barat<br/>Nomor&nbsp;&nbsp;: __________________<br/>Tanggal: __________________", attachment_style), Spacer(1, 2 * mm), Paragraph(report_data["title"].replace("\n", "<br/>"), title_style), Paragraph(report_data["period_label"], title_style), Paragraph(report_data["period_text"], title_style), Spacer(1, 2 * mm)]
    all_months = list(range(1, 13))
    header_width = 18
    rows = [
        ["NO", "UNIT KERJA", "NAMA", "NIP", "JUMLAH KETIDAKHADIRAN MASUK KERJA TANPA KETERANGAN (TK)"] + [""] * 11 + ["TOTAL", "KETERANGAN"],
        ["", "", "", "", "TRIWULAN I", "", "", "TRIWULAN II", "", "", "TRIWULAN III", "", "", "TRIWULAN IV", "", "", "", ""],
        ["", "", "", ""] + MONTH_SHORT + ["", ""],
    ]
    for opd_number, opd in enumerate(report_data["opds"], 1):
        rows.append([str(opd_number), f"{opd['opd_name'].upper()} (TOTAL TK: {opd['total_tk']})"] + [""] * (header_width - 2))
        for employee_number, employee in enumerate(opd["employees"], 1):
            rows.append([str(employee_number), "", employee["nama"], employee["nip"]] + [str(employee["monthly_tk"][m] or "-") for m in all_months] + [str(employee["total_tk"]), employee["keterangan"]])
    widths = [7 * mm, 39 * mm, 39 * mm, 31 * mm] + [7.5 * mm] * 12 + [10 * mm, 18 * mm]
    table = LongTable(rows, colWidths=widths, repeatRows=3, splitByRow=1)
    style = TableStyle([
        ("FONTNAME", (0, 0), (-1, 2), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 6.5),
        ("BACKGROUND", (0, 0), (-1, 2), colors.HexColor("#EAF0F7")), ("TEXTCOLOR", (0, 0), (-1, 2), colors.HexColor("#173B63")),
        ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#808080")), ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("ALIGN", (1, 1), (3, -1), "LEFT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ])
    for column in [0, 1, 2, 3, 16, 17]: style.add("SPAN", (column, 0), (column, 2))
    style.add("SPAN", (4, 0), (15, 0))
    for start in [4, 7, 10, 13]: style.add("SPAN", (start, 1), (start + 2, 1))
    row_index = 3
    for opd in report_data["opds"]:
        style.add("SPAN", (1, row_index), (-1, row_index)); style.add("BACKGROUND", (0, row_index), (-1, row_index), colors.HexColor("#F3F4F6")); style.add("FONTNAME", (0, row_index), (-1, row_index), "Helvetica-Bold")
        row_index += len(opd["employees"]) + 1
    table.setStyle(style); story.append(table); document.build(story)
    return output.getvalue()
